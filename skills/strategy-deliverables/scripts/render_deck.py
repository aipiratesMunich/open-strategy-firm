#!/usr/bin/env python3
"""Render a Markdown deck spec (06-deck.md) into a self-contained HTML deck
and, if python-pptx is installed, into a .pptx.

The HTML deck needs no network: inline CSS, inline SVG charts, keyboard
navigation (arrows, space, Home/End, t = titles-only read-through,
n = speaker notes) and one slide per page when printed to PDF.

Chart types drawn from the slide's data table: bar, hbar, line, stacked,
stacked-hbar, waterfall. Every run also prints spec warnings (missing
source, label-like titles, chart that does not fit the message type,
waterfall totals that do not add up).

Standard library only. python-pptx is optional (pip install python-pptx).

Usage:
    python3 render_deck.py engagement/acme/06-deck.md
    python3 render_deck.py 06-deck.md --out-dir build --no-pptx
    python3 render_deck.py 06-deck.md --titles      # titles-only read-through
    python3 render_deck.py 06-deck.md --check       # warnings only, no files
"""

import argparse
import html
import math
import re
import sys
from pathlib import Path

DEFAULT_THEME = {
    "text": "#14213D",
    "accent": "#D9822B",
    "muted": "#B8C0CC",
    "light": "#E8EBF0",
    "background": "#FFFFFF",
    "font": "'Helvetica Neue', Helvetica, Arial, sans-serif",
}
SLIDE_RE = re.compile(r"^##\s+([A-Za-z]?\d+[a-z]?)\.\s+(.+?)\s*$")
KEYS = {
    "layout", "tracker", "message", "chart", "header", "highlight", "reference",
    "takeaway", "source", "note", "lead", "ask", "subtitle", "footnote", "scale", "body",
}
LAYOUTS = {"title", "summary", "section", "chart", "table", "text"}
CHARTS = {"bar", "hbar", "line", "stacked", "stacked-hbar", "waterfall"}
MESSAGE_CHARTS = {
    "ranking": {"hbar", "bar"},
    "time-series": {"bar", "line"},
    "part-to-whole": {"stacked", "stacked-hbar"},
    "distribution": {"bar"},
    "deviation": {"bar", "hbar"},
    "flow": {"waterfall"},
    "correlation": set(),
}
SW, SH = 1280, 720


# ---------------------------------------------------------------- parsing

def parse_spec(text):
    lines = text.splitlines()
    meta, theme = {}, dict(DEFAULT_THEME)
    i = 0
    if lines and lines[0].strip() == "---":
        i, in_theme = 1, False
        while i < len(lines) and lines[i].strip() != "---":
            ln = lines[i]
            i += 1
            m = re.match(r"^(\s*)([\w-]+):\s*(.*)$", ln)
            if not m or ln.lstrip().startswith("#"):
                continue
            indent, key, val = m.group(1), m.group(2), m.group(3).strip().strip("\"'")
            if indent and in_theme:
                theme[key] = val
            elif key == "theme" and not val:
                in_theme = True
            else:
                in_theme = False
                meta[key] = val
        i += 1
    slides, cur, last = [], None, None
    for ln in lines[i:]:
        m = SLIDE_RE.match(ln)
        if m:
            cur = {"id": m.group(1), "title": m.group(2).strip(), "f": {}, "body": [], "table": []}
            slides.append(cur)
            last = None
            continue
        if ln.startswith("#"):
            cur, last = None, None  # any other heading ends the slide block
            continue
        if cur is None:
            continue
        s = ln.strip()
        if not s:
            continue
        if s.startswith("|"):
            cells = [c.strip() for c in s.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                continue
            cur["table"].append(cells)
            last = None
            continue
        mb = re.match(r"^(?:[-*]|\d+[.)])\s+(.*)$", s)
        if mb:
            cur["body"].append(mb.group(1))
            last = "body"
            continue
        mk = re.match(r"^([a-z][\w-]*):\s*(.*)$", s)
        if mk and mk.group(1) in KEYS:
            last = mk.group(1)
            if last != "body":
                cur["f"][last] = mk.group(2).strip()
            continue
        if last == "body" and cur["body"]:
            cur["body"][-1] += " " + s
        elif last:
            cur["f"][last] = (cur["f"].get(last, "") + " " + s).strip()
    for s in slides:
        f = s["f"]
        if "layout" not in f:
            f["layout"] = "chart" if f.get("chart") else ("table" if s["table"] else "text")
        f["layout"] = f["layout"].lower()
        if "chart" in f:
            f["chart"] = f["chart"].lower()
    return meta, theme, slides


def num(cell):
    if cell is None:
        return None
    s = re.sub(r"(?i)eur|usd|chf|gbp|[€$£%\s]", "", str(cell)).replace(",", "").replace("−", "-")
    s = s.lstrip("+")
    if s.endswith(("m", "k", "x")):
        s = s[:-1]
    try:
        return float(s)
    except ValueError:
        return None


def chart_data(slide):
    t = slide["table"]
    if len(t) < 2:
        return None
    head, rows = t[0], t[1:]
    kind_col = None
    for j, h in enumerate(head):
        if j > 0 and h.lower() in ("type", "kind"):
            kind_col = j
    ser_cols = [j for j in range(1, len(head)) if j != kind_col]
    cats = [r[0] if r else "" for r in rows]
    series = []
    for j in ser_cols:
        raw = [r[j] if j < len(r) else "" for r in rows]
        series.append({"name": head[j], "raw": raw, "vals": [num(x) for x in raw]})
    kinds = None
    if kind_col is not None:
        kinds = [(r[kind_col].lower() if kind_col < len(r) else "delta") for r in rows]
    return {"head": head, "cats": cats, "series": series, "kinds": kinds}


def highlights(slide):
    out = []
    for h in slide["f"].get("highlight", "").split(";"):
        h = h.strip()
        if h:
            out.append(tuple(p.strip().lower() for p in h.split("/", 1)))
    return out


def is_hl(hls, cat, ser):
    cat, ser = (cat or "").lower(), (ser or "").lower()
    for h in hls:
        if len(h) == 2 and h[0] == cat and h[1] == ser:
            return True
        if len(h) == 1 and h[0] in (cat, ser):
            return True
    return False


def reference(slide):
    ref = slide["f"].get("reference", "")
    if not ref:
        return None
    val, _, label = ref.partition("|")
    v = num(val)
    return None if v is None else (v, label.strip())


def waterfall_steps(data):
    s = data["series"][0]
    n = len(data["cats"])
    kinds = data["kinds"] or ["total" if k in (0, n - 1) and n > 2 else "delta" for k in range(n)]
    steps, running, problems = [], 0.0, []
    for k in range(n):
        v, raw, kind = s["vals"][k], s["raw"][k], kinds[k]
        if kind.startswith("total") or kind in ("start", "end", "subtotal"):
            if v is None:
                v, raw = running, fmt(running)
            elif k > 0 and abs(v - running) > max(0.011 * abs(running), 0.05):
                problems.append("total '{}' is {} but steps add up to {}".format(data["cats"][k], fmt(v), fmt(running)))
            running = v
            steps.append({"cat": data["cats"][k], "lo": min(0, v), "hi": max(0, v), "raw": raw, "total": True})
        else:
            v = v or 0.0
            start = running
            running += v
            lab = raw if raw.strip().startswith(("+", "-", "−")) or v < 0 else "+" + raw
            steps.append({"cat": data["cats"][k], "lo": min(start, running), "hi": max(start, running), "raw": lab, "total": False, "neg": v < 0})
    return steps, problems


def lint_slides(meta, slides):
    warns = []
    main = [s for s in slides if not s["id"][0].isalpha()]
    if len(main) > 15:
        warns.append("deck: {} main slides; a board deck is usually 8 to 15, move detail to the appendix".format(len(main)))
    for s in slides:
        f, lay, sid = s["f"], s["f"]["layout"], s["id"]
        appendix = sid[0].isalpha()
        if lay not in LAYOUTS:
            warns.append("slide {}: unknown layout '{}'".format(sid, lay))
        if len(s["title"]) > 120 and lay not in ("summary",):
            warns.append("slide {}: title has {} characters, likely more than two lines".format(sid, len(s["title"])))
        if lay in ("chart", "table", "text") and not appendix and len(s["title"].split()) < 5:
            warns.append("slide {}: title '{}' looks like a label, write the finding as a sentence".format(sid, s["title"]))
        if lay in ("chart", "table") and not f.get("source"):
            warns.append("slide {}: no source line".format(sid))
        if lay == "chart":
            ch = f.get("chart", "")
            if ch not in CHARTS:
                warns.append("slide {}: chart '{}' not supported ({})".format(sid, ch, ", ".join(sorted(CHARTS))))
                continue
            data = chart_data(s)
            if not data or any(v is None for ser in data["series"] for v in ser["vals"]) and ch != "waterfall":
                warns.append("slide {}: chart needs a table with a label column and numeric value columns".format(sid))
                continue
            msg = f.get("message", "").lower()
            if msg in MESSAGE_CHARTS and ch not in MESSAGE_CHARTS[msg]:
                warns.append("slide {}: message '{}' is usually shown as {}, not {}".format(
                    sid, msg, " or ".join(sorted(MESSAGE_CHARTS[msg])) or "a scatter plot (draw it elsewhere)", ch))
            if ch == "waterfall":
                for p in waterfall_steps(data)[1]:
                    warns.append("slide {}: waterfall {}".format(sid, p))
            if not f.get("header"):
                warns.append("slide {}: no chart header (measure, unit, period)".format(sid))
    return warns


# ---------------------------------------------------------------- helpers

def esc(s):
    return html.escape(str(s or ""), quote=True)


def inline(s):
    out = esc(s)
    out = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", out)
    return re.sub(r"(?<!\*)\*(?!\*)(.+?)\*", r"<em>\1</em>", out)


def fmt(v):
    if v is None:
        return ""
    if abs(v - round(v)) < 1e-9:
        return "{:,}".format(int(round(v)))
    return ("{:,.2f}".format(v)).rstrip("0").rstrip(".")


def nice(v):
    if v <= 0:
        return 1.0
    e = 10 ** math.floor(math.log10(v))
    for n in (1, 1.2, 1.5, 2, 2.5, 3, 4, 5, 6, 8, 10):
        if v / e <= n + 1e-9:
            return n * e
    return 10 * e


def hex_rgb(c):
    c = c.lstrip("#")
    if len(c) == 3:
        c = "".join(ch * 2 for ch in c)
    return tuple(int(c[i:i + 2], 16) for i in (0, 2, 4))


def mix(a, b, t):
    ra, rb = hex_rgb(a), hex_rgb(b)
    return "#%02X%02X%02X" % tuple(int(ra[i] + (rb[i] - ra[i]) * t) for i in range(3))


def wrap(text, width_px, size=14):
    max_chars = max(4, int(width_px / (size * 0.55)))
    words, lines, cur = str(text).split(), [], ""
    for w in words:
        if cur and len(cur) + 1 + len(w) > max_chars:
            lines.append(cur)
            cur = w
        else:
            cur = (cur + " " + w).strip()
    if cur:
        lines.append(cur)
    return lines[:3]


def series_colors(theme, n, hl_any):
    if hl_any:
        return [theme["muted"]] * n if n == 1 else [mix(theme["muted"], theme["light"], k / max(1, n)) for k in range(n)]
    if n == 1:
        return [theme["text"]]
    return [mix(theme["text"], theme["light"], 0.75 * k / (n - 1)) for k in range(n)]


def text_on(fill):
    r, g, b = hex_rgb(fill)
    return "#FFFFFF" if (0.299 * r + 0.587 * g + 0.114 * b) < 150 else "#14213D"


def svg_text(x, y, s, size=14, anchor="middle", fill="#14213D", weight="normal", halo=True):
    extra = ' stroke="#FFFFFF" stroke-width="4" paint-order="stroke" stroke-linejoin="round"' if halo else ""
    return '<text x="{:.1f}" y="{:.1f}" font-size="{}" text-anchor="{}" fill="{}" font-weight="{}"{}>{}</text>'.format(
        x, y, size, anchor, fill, weight, extra, esc(s))


def legend(series, colors, theme, y=16):
    parts, x = [], 0
    for ser, col in zip(series, colors):
        parts.append('<rect x="{}" y="{}" width="12" height="12" fill="{}"/>'.format(x, y - 11, col))
        parts.append(svg_text(x + 18, y, ser["name"], 14, "start", theme["text"]))
        x += 40 + len(ser["name"]) * 8
    return "".join(parts)


# ---------------------------------------------------------------- SVG charts

def svg_bar(slide, data, theme, W, H):
    hls, ref = highlights(slide), reference(slide)
    cats, series = data["cats"], data["series"]
    k, n = len(series), len(cats)
    cols = series_colors(theme, k, bool(hls))
    vals = [v for s in series for v in s["vals"]] + ([ref[0]] if ref else [])
    vmax, vmin = nice(max(vals + [0])), min(vals + [0])
    vmin = -nice(-vmin) if vmin < 0 else 0
    top, bottom, left, right = 44 if k > 1 else 28, 52, 8, 8 + (150 if ref else 0)
    ph, pw = H - top - bottom, W - left - right
    y = lambda v: top + ph * (vmax - v) / (vmax - vmin)
    slot = pw / n
    bw = min(slot * 0.62 / k, 120)
    out = [legend(series, cols, theme) if k > 1 else ""]
    for i, cat in enumerate(cats):
        gx = left + slot * i + (slot - bw * k) / 2
        for j, s in enumerate(series):
            v = s["vals"][i]
            col = theme["accent"] if is_hl(hls, cat, s["name"]) else cols[j]
            y0, y1 = sorted((y(0), y(v)))
            out.append('<rect x="{:.1f}" y="{:.1f}" width="{:.1f}" height="{:.1f}" fill="{}"/>'.format(gx + bw * j, y0, bw - 4, max(y1 - y0, 0.5), col))
            ly = y(v) - 8 if v >= 0 else y(v) + 18
            out.append(svg_text(gx + bw * j + (bw - 4) / 2, ly, s["raw"][i], 15, fill=theme["text"], weight="bold" if is_hl(hls, cat, s["name"]) else "normal"))
        for li, line in enumerate(wrap(cat, slot - 8)):
            out.append(svg_text(left + slot * i + slot / 2, H - bottom + 22 + li * 17, line, 14, fill=theme["text"]))
    out.append('<line x1="{}" x2="{}" y1="{:.1f}" y2="{:.1f}" stroke="{}" stroke-width="1"/>'.format(left, left + pw, y(0), y(0), theme["text"]))
    if ref:
        out.insert(1, ref_line(ref, y(ref[0]), left, left + pw, theme))
    return out


def ref_line(ref, yy, x0, x1, theme):
    return ('<line x1="{:.1f}" x2="{:.1f}" y1="{:.1f}" y2="{:.1f}" stroke="{}" stroke-width="1.5" stroke-dasharray="6 5"/>'.format(x0, x1, yy, yy, theme["text"])
            + svg_text(x1 + 8, yy + 5, ref[1] or fmt(ref[0]), 14, "start", theme["text"], "bold"))


def svg_hbar(slide, data, theme, W, H):
    hls, ref = highlights(slide), reference(slide)
    cats, series = data["cats"], data["series"]
    k, n = len(series), len(cats)
    cols = series_colors(theme, k, bool(hls))
    vals = [v for s in series for v in s["vals"]] + ([ref[0]] if ref else [])
    vmax = nice(max(vals + [0]))
    lab_w = min(280, max(len(c) for c in cats) * 8 + 16)
    top = 36 if k > 1 else 8
    left, right = lab_w, 70
    pw, ph = W - left - right, H - top - 8
    slot = ph / n
    bh = min(slot * 0.66 / k, 46)
    x = lambda v: left + pw * max(v, 0) / vmax
    out = [legend(series, cols, theme) if k > 1 else ""]
    for i, cat in enumerate(cats):
        gy = top + slot * i + (slot - bh * k) / 2
        for j, s in enumerate(series):
            v = s["vals"][i]
            hl = is_hl(hls, cat, s["name"])
            out.append('<rect x="{}" y="{:.1f}" width="{:.1f}" height="{:.1f}" fill="{}"/>'.format(left, gy + bh * j, max(x(v) - left, 0.5), bh - 3, theme["accent"] if hl else cols[j]))
            out.append(svg_text(x(v) + 8, gy + bh * j + bh / 2 + 4, s["raw"][i], 15, "start", theme["text"], "bold" if hl else "normal"))
        out.append(svg_text(left - 12, top + slot * i + slot / 2 + 5, cat, 15, "end", theme["text"], "bold" if is_hl(hls, cat, "") else "normal"))
    out.append('<line x1="{0}" x2="{0}" y1="{1}" y2="{2:.1f}" stroke="{3}"/>'.format(left, top, top + ph, theme["text"]))
    if ref:
        xx = x(ref[0])
        out.append('<line x1="{0:.1f}" x2="{0:.1f}" y1="{1}" y2="{2:.1f}" stroke="{3}" stroke-dasharray="6 5"/>'.format(xx, top, top + ph, theme["text"]))
        out.append(svg_text(xx, H - 2, ref[1] or fmt(ref[0]), 13, "middle", theme["text"], "bold"))
    return out


def svg_line(slide, data, theme, W, H):
    hls, ref = highlights(slide), reference(slide)
    cats, series = data["cats"], data["series"]
    k, n = len(series), len(cats)
    cols = series_colors(theme, k, bool(hls))
    vals = [v for s in series for v in s["vals"]] + ([ref[0]] if ref else [])
    if slide["f"].get("scale", "zero").lower() == "fit":
        span = max(vals) - min(vals) or 1
        vmin, vmax = min(vals) - span * 0.15, max(vals) + span * 0.15
    else:
        vmin, vmax = min(0, min(vals)), nice(max(vals))
    top, bottom, left, right = 30, 40, 56, 190
    pw, ph = W - left - right, H - top - bottom
    x = lambda i: left + (pw * i / (n - 1) if n > 1 else pw / 2)
    y = lambda v: top + ph * (vmax - v) / (vmax - vmin)
    out = []
    for t in (vmin, (vmin + vmax) / 2, vmax):
        out.append(svg_text(left - 10, y(t) + 5, fmt(round(t, 1)), 13, "end", theme["muted"]))
    out.append('<line x1="{}" x2="{}" y1="{:.1f}" y2="{:.1f}" stroke="{}"/>'.format(left, left + pw, y(vmin), y(vmin), theme["muted"]))
    for i, c in enumerate(cats):
        out.append(svg_text(x(i), H - bottom + 24, c, 14, fill=theme["text"]))
    if ref:
        out.append(ref_line(ref, y(ref[0]), left, left + pw, theme))
    for j, s in enumerate(series):
        hl = is_hl(hls, "", s["name"]) or (k == 1 and not hls)
        col = theme["accent"] if is_hl(hls, "", s["name"]) else cols[j]
        pts = " ".join("{:.1f},{:.1f}".format(x(i), y(v)) for i, v in enumerate(s["vals"]))
        out.append('<polyline points="{}" fill="none" stroke="{}" stroke-width="{}"/>'.format(pts, col, 3.5 if hl else 2.5))
        for i, v in enumerate(s["vals"]):
            pt_hl = is_hl(hls, cats[i], s["name"]) and len([h for h in hls if len(h) == 2]) > 0
            out.append('<circle cx="{:.1f}" cy="{:.1f}" r="{}" fill="{}"/>'.format(x(i), y(v), 6 if pt_hl else 4, theme["accent"] if pt_hl else col))
            if hl or pt_hl or i == n - 1:
                out.append(svg_text(x(i), y(v) - 12, s["raw"][i], 14, fill=theme["text"], weight="bold" if pt_hl else "normal"))
        if k > 1:
            out.append(svg_text(x(n - 1) + 12, y(s["vals"][-1]) + 5, s["name"], 14, "start", col, "bold"))
    return out


def svg_stacked(slide, data, theme, W, H, horizontal):
    hls = highlights(slide)
    cats, series = data["cats"], data["series"]
    k, n = len(series), len(cats)
    cols = series_colors(theme, k, bool(hls))
    totals = [sum(s["vals"][i] or 0 for s in series) for i in range(n)]
    vmax = nice(max(totals))
    single = horizontal and n == 1
    out = [] if single else [legend(series, cols, theme)]
    top = 40
    if horizontal:
        lab_w = min(280, max(len(c) for c in cats) * 8 + 16)
        left, right = lab_w, 80
        pw, ph = W - left - right, H - top - 4
        slot = ph / n
        bh = min(slot * 0.66, 70 if not single else 90)
        for i, cat in enumerate(cats):
            yy, xx = top + slot * i + (slot - bh) / 2, left
            for j, s in enumerate(series):
                v = s["vals"][i] or 0
                w = pw * v / vmax
                col = theme["accent"] if is_hl(hls, cat, s["name"]) else cols[j]
                out.append('<rect x="{:.1f}" y="{:.1f}" width="{:.1f}" height="{:.1f}" fill="{}" stroke="{}" stroke-width="1"/>'.format(xx, yy, w, bh, col, theme["background"]))
                if w > len(s["raw"][i]) * 9 + 8:
                    out.append(svg_text(xx + w / 2, yy + bh / 2 + 5, s["raw"][i], 14, fill=text_on(col), halo=False))
                if single:
                    for li, line in enumerate(wrap(s["name"], w - 6)):
                        out.append(svg_text(xx + w / 2, yy - 12 - 17 * (len(wrap(s["name"], w - 6)) - 1 - li), line, 14, fill=theme["text"], weight="bold" if is_hl(hls, cat, s["name"]) else "normal"))
                xx += w
            out.append(svg_text(left - 12, yy + bh / 2 + 5, cat, 15, "end", theme["text"]))
            out.append(svg_text(xx + 8, yy + bh / 2 + 5, fmt(totals[i]), 14, "start", theme["text"], "bold"))
    else:
        bottom, left = 52, 8
        pw, ph = W - 16, H - top - bottom - 20
        slot = pw / n
        bw = min(slot * 0.6, 140)
        for i, cat in enumerate(cats):
            xx, base = left + slot * i + (slot - bw) / 2, top + 20 + ph
            for j, s in enumerate(series):
                v = s["vals"][i] or 0
                h = ph * v / vmax
                col = theme["accent"] if is_hl(hls, cat, s["name"]) else cols[j]
                out.append('<rect x="{:.1f}" y="{:.1f}" width="{:.1f}" height="{:.1f}" fill="{}" stroke="{}" stroke-width="1"/>'.format(xx, base - h, bw, h, col, theme["background"]))
                if h > 20:
                    out.append(svg_text(xx + bw / 2, base - h / 2 + 5, s["raw"][i], 14, fill=text_on(col), halo=False))
                base -= h
            out.append(svg_text(xx + bw / 2, base - 8, fmt(totals[i]), 14, fill=theme["text"], weight="bold"))
            for li, line in enumerate(wrap(cat, slot - 8)):
                out.append(svg_text(left + slot * i + slot / 2, H - bottom + 22 + li * 17, line, 14, fill=theme["text"]))
    return out


def svg_waterfall(slide, data, theme, W, H):
    hls, ref = highlights(slide), reference(slide)
    steps, _ = waterfall_steps(data)
    vals = [s["hi"] for s in steps] + [s["lo"] for s in steps] + ([ref[0]] if ref else [])
    vmax, vmin = nice(max(vals)), min(0, min(vals))
    top, bottom, left, right = 28, 52, 8, 8 + (150 if ref else 0)
    pw, ph = W - left - right, H - top - bottom
    y = lambda v: top + ph * (vmax - v) / (vmax - vmin)
    n = len(steps)
    slot = pw / n
    bw = min(slot * 0.62, 120)
    out = []
    prev_end = None
    for i, s in enumerate(steps):
        xx = left + slot * i + (slot - bw) / 2
        if s["total"]:
            col = theme["accent"] if is_hl(hls, s["cat"], "") else theme["text"]
        else:
            col = theme["accent"] if is_hl(hls, s["cat"], "") else theme["muted"]
        out.append('<rect x="{:.1f}" y="{:.1f}" width="{:.1f}" height="{:.1f}" fill="{}"/>'.format(xx, y(s["hi"]), bw, max(y(s["lo"]) - y(s["hi"]), 0.5), col))
        out.append(svg_text(xx + bw / 2, y(s["hi"]) - 8, s["raw"], 15, fill=theme["text"], weight="bold" if is_hl(hls, s["cat"], "") or s["total"] else "normal"))
        level = s["hi"] if (s["total"] or not s.get("neg")) else s["lo"]
        if prev_end is not None:
            out.append('<line x1="{:.1f}" x2="{:.1f}" y1="{:.1f}" y2="{:.1f}" stroke="{}" stroke-dasharray="3 3"/>'.format(prev_end[0], xx, prev_end[1], prev_end[1], theme["muted"]))
        prev_end = (xx + bw, y(level))
        for li, line in enumerate(wrap(s["cat"], slot - 8)):
            out.append(svg_text(left + slot * i + slot / 2, H - bottom + 22 + li * 17, line, 14, fill=theme["text"]))
    out.append('<line x1="{}" x2="{}" y1="{:.1f}" y2="{:.1f}" stroke="{}"/>'.format(left, left + pw, y(0), y(0), theme["text"]))
    if ref:
        out.insert(1, ref_line(ref, y(ref[0]), left, left + pw, theme))
    return out


def render_svg(slide, theme, W, H):
    data = chart_data(slide)
    ch = slide["f"].get("chart", "")
    if not data or ch not in CHARTS:
        return '<p class="warn">Chart could not be drawn: check chart type and data table.</p>'
    if ch == "bar":
        parts = svg_bar(slide, data, theme, W, H)
    elif ch == "hbar":
        parts = svg_hbar(slide, data, theme, W, H)
    elif ch == "line":
        parts = svg_line(slide, data, theme, W, H)
    elif ch == "waterfall":
        parts = svg_waterfall(slide, data, theme, W, H)
    else:
        parts = svg_stacked(slide, data, theme, W, H, ch == "stacked-hbar")
    return '<svg class="chart" viewBox="0 0 {} {}" preserveAspectRatio="xMidYMid meet" role="img" aria-label="{}">{}</svg>'.format(
        W, H, esc(slide["title"]), "".join(parts))


# ---------------------------------------------------------------- HTML

def html_table(slide, theme):
    t = slide["table"]
    if not t:
        return ""
    hls = highlights(slide)
    numeric = [j > 0 and all(num(r[j]) is not None for r in t[1:] if j < len(r)) for j in range(len(t[0]))]
    th = lambda j, c: '<th{}>{}</th>'.format(' class="num"' if j < len(numeric) and numeric[j] else "", inline(c))
    rows = ["<tr>" + "".join(th(j, c) for j, c in enumerate(t[0])) + "</tr>"]
    for r in t[1:]:
        cls = ' class="hl"' if r and is_hl(hls, r[0], "") else ""
        rows.append("<tr{}>".format(cls) + "".join(
            '<td{}>{}</td>'.format(' class="num"' if j < len(numeric) and numeric[j] else "", inline(c)) for j, c in enumerate(r)) + "</tr>")
    return '<table class="data">{}</table>'.format("".join(rows))


def bullets(items, ordered=False):
    if not items:
        return ""
    tag = "ol" if ordered else "ul"
    return "<{0}>{1}</{0}>".format(tag, "".join("<li>{}</li>".format(inline(b)) for b in items))


def slide_html(s, idx, meta, theme):
    f, lay = s["f"], s["f"]["layout"]
    appendix = s["id"][0].isalpha()
    tracker = f.get("tracker") or ("Appendix" if appendix else "")
    parts = []
    if lay == "title":
        parts.append('<div class="t-client">{}</div>'.format(inline(meta.get("client", ""))))
        parts.append('<h1 class="t-title">{}</h1>'.format(inline(s["title"])))
        if f.get("subtitle"):
            parts.append('<p class="t-sub">{}</p>'.format(inline(f["subtitle"])))
        parts.append('<p class="t-meta">{}</p>'.format(inline(" | ".join(x for x in (meta.get("audience"), meta.get("date")) if x))))
    elif lay == "section":
        parts.append('<div class="tracker">{}</div><h2 class="s-title">{}</h2>'.format(inline(tracker), inline(s["title"])))
        if f.get("lead"):
            parts.append('<p class="lead">{}</p>'.format(inline(f["lead"])))
    else:
        parts.append('<div class="tracker">{}</div>'.format(inline(tracker)))
        parts.append('<h2 class="title{}">{}</h2>'.format(" long" if len(s["title"]) > 95 else "", inline(s["title"])))
        if f.get("header"):
            parts.append('<div class="header">{}</div>'.format(inline(f["header"])))
        if lay == "summary":
            body = ""
            if f.get("lead"):
                body += '<p class="lead">{}</p>'.format(inline(f["lead"]))
            body += bullets(s["body"], ordered=True)
            if f.get("ask"):
                body += '<div class="ask">{}</div>'.format(inline(f["ask"]))
            parts.append('<div class="content summary">{}</div>'.format(body))
        elif lay == "chart":
            two = bool(s["body"])
            svg = render_svg(s, theme, 700 if two else 1150, 400 if not f.get("takeaway") else 360)
            side = '<div class="side">{}</div>'.format(bullets(s["body"])) if two else ""
            parts.append('<div class="content{}"><div class="exhibit">{}</div>{}</div>'.format(" two" if two else "", svg, side))
        elif lay == "table":
            side = '<div class="side">{}</div>'.format(bullets(s["body"])) if s["body"] else ""
            parts.append('<div class="content{}"><div class="exhibit tbl">{}</div>{}</div>'.format(" two" if side else "", html_table(s, theme), side))
        else:
            body = ('<p class="lead">{}</p>'.format(inline(f["lead"])) if f.get("lead") else "") + bullets(s["body"])
            parts.append('<div class="content text">{}</div>'.format(body))
        if f.get("takeaway"):
            parts.append('<div class="takeaway">{}</div>'.format(inline(f["takeaway"])))
    foot = []
    if f.get("footnote"):
        foot.append('<div class="footnote">{}</div>'.format(inline(f["footnote"])))
    if f.get("source"):
        foot.append('<div class="source">Source: {}</div>'.format(inline(f["source"])))
    parts.append('<footer>{}<span class="pn">{}</span></footer>'.format("".join(foot) or "<div></div>", esc(s["id"])))
    if f.get("note"):
        parts.append('<aside class="notes"><b>Speaker note</b> {}</aside>'.format(inline(f["note"])))
    return '<section class="slide l-{}" id="s{}" data-title="{}">{}</section>'.format(lay, idx + 1, esc(s["title"]), "".join(parts))


CSS = """
:root{--text:%(text)s;--accent:%(accent)s;--muted:%(muted)s;--light:%(light)s;--bg:%(background)s}
*{box-sizing:border-box;margin:0;padding:0}
html,body{height:100%%;background:#2B2F36;font-family:%(font)s;color:var(--text)}
#stage{position:absolute;left:0;top:0;width:1280px;height:720px;transform-origin:0 0}
.slide{display:none;position:absolute;inset:0;width:1280px;height:720px;background:var(--bg);padding:44px 64px 30px;flex-direction:column}
.slide.active{display:flex}
.tracker{font-size:14px;letter-spacing:.08em;text-transform:uppercase;color:var(--accent);font-weight:700;min-height:18px;margin-bottom:10px}
.title{font-size:30px;line-height:1.22;font-weight:700;max-width:1120px}
.title.long{font-size:26px}
.header{font-size:16px;color:#5B6475;margin-top:12px}
.content{flex:1;display:flex;gap:40px;margin-top:22px;min-height:0}
.exhibit{flex:1;min-width:0;display:flex;align-items:center;justify-content:center}
.exhibit.tbl{align-items:flex-start}
.chart{width:100%%;height:100%%}
.side{width:380px;font-size:19px;line-height:1.45;padding-top:8px}
.side li,.text li{margin:0 0 14px 20px}
.text{flex-direction:column;font-size:22px;line-height:1.45;max-width:1040px}
.lead{font-size:20px;line-height:1.5;color:#3B4458;margin-bottom:18px;max-width:1080px}
.summary{flex-direction:column;gap:0}
.summary ol{font-size:19px;line-height:1.45;margin-left:24px}
.summary li{margin-bottom:10px;padding-left:6px}
.ask{margin-top:auto;border-left:5px solid var(--accent);background:var(--light);padding:14px 18px;font-size:18px;line-height:1.4}
.takeaway{margin-top:14px;padding:10px 16px;border-top:2px solid var(--text);font-size:18px;font-weight:600}
footer{display:flex;justify-content:space-between;align-items:flex-end;gap:24px;margin-top:14px;font-size:13px;color:#5B6475;line-height:1.35}
footer .pn{font-weight:700;color:var(--text)}
.footnote{margin-bottom:2px}
table.data{border-collapse:collapse;font-size:17px;width:100%%}
table.data th{text-align:left;border-bottom:2px solid var(--text);padding:8px 12px;font-weight:700}
table.data td{border-bottom:1px solid var(--light);padding:8px 12px;vertical-align:top}
table.data th.num{text-align:right}
table.data td.num{text-align:right;font-variant-numeric:tabular-nums}
table.data tr.hl td{background:var(--light);font-weight:700}
.slide.l-title{justify-content:center;padding-left:96px}
.t-client{font-size:18px;letter-spacing:.08em;text-transform:uppercase;color:var(--accent);font-weight:700;margin-bottom:24px}
.t-title{font-size:48px;line-height:1.15;max-width:980px;font-weight:700}
.t-sub{font-size:22px;margin-top:24px;color:#3B4458;max-width:900px}
.t-meta{font-size:16px;margin-top:40px;color:#5B6475}
.slide.l-title footer,.slide.l-section footer{margin-top:auto}
.slide.l-section{justify-content:center;padding-left:96px}
.s-title{font-size:40px;line-height:1.2;max-width:1000px;margin-top:12px}
.notes{display:none;position:absolute;left:0;right:0;bottom:0;background:#FFF8E8;border-top:2px solid var(--accent);padding:12px 64px;font-size:16px;line-height:1.4}
body.shownotes .slide.active .notes{display:block}
#overview{display:none;position:fixed;inset:0;overflow:auto;background:var(--bg);padding:40px 56px;font-size:18px;line-height:1.5}
body.overview #overview{display:block}
#overview h1{font-size:22px;margin-bottom:16px}
#overview li{margin:0 0 8px 28px;cursor:pointer}
#hint{position:fixed;right:12px;bottom:8px;font-size:12px;color:#C9CED8}
.warn{color:#B00020}
@media print{
  @page{size:1280px 720px;margin:0}
  html,body{background:#fff}
  #stage{position:static;transform:none!important;width:auto;height:auto}
  .slide{display:flex!important;position:relative;page-break-after:always;break-after:page}
  .notes,#hint,#overview{display:none!important}
}
"""

JS = """
(function(){
var slides=[].slice.call(document.querySelectorAll('.slide')),i=0,stage=document.getElementById('stage');
function fit(){var s=Math.min(innerWidth/1280,innerHeight/720);stage.style.transform='translate('+((innerWidth-1280*s)/2)+'px,'+((innerHeight-720*s)/2)+'px) scale('+s+')';}
function show(n){i=Math.max(0,Math.min(slides.length-1,n));slides.forEach(function(el,k){el.classList.toggle('active',k===i);});history.replaceState(null,'','#'+(i+1));}
addEventListener('resize',fit);
addEventListener('keydown',function(e){
 var k=e.key;
 if(k==='ArrowRight'||k==='PageDown'||k===' '||k==='Enter'){show(i+1);e.preventDefault();}
 else if(k==='ArrowLeft'||k==='PageUp'||k==='Backspace'){show(i-1);e.preventDefault();}
 else if(k==='Home'){show(0);}else if(k==='End'){show(slides.length-1);}
 else if(k==='t'){document.body.classList.toggle('overview');}
 else if(k==='n'){document.body.classList.toggle('shownotes');}
 else if(k==='Escape'){document.body.classList.remove('overview');}
});
[].slice.call(document.querySelectorAll('#overview li')).forEach(function(li,k){li.addEventListener('click',function(){document.body.classList.remove('overview');show(k);});});
fit();show((parseInt(location.hash.slice(1),10)||1)-1);
})();
"""


def build_html(meta, theme, slides):
    secs = "".join(slide_html(s, k, meta, theme) for k, s in enumerate(slides))
    overview = "".join("<li>{} <span style='color:#8A93A3'>({})</span></li>".format(inline(s["title"]), esc(s["id"])) for s in slides)
    return """<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{title}</title><style>{css}</style></head>
<body><div id="stage">{secs}</div>
<div id="overview"><h1>Titles-only read-through</h1><ol>{ov}</ol></div>
<div id="hint">arrows: navigate | t: titles | n: notes | print: PDF</div>
<script>{js}</script></body></html>
""".format(title=esc(meta.get("title", "Deck")), css=CSS % theme, secs=secs, ov=overview, js=JS)


def build_memo_html(meta, theme, text):
    body_md = text.split("---", 2)[2] if text.startswith("---") else text
    out, in_list, table = [], None, []

    def flush():
        nonlocal in_list, table
        if in_list:
            out.append("</{}>".format(in_list))
            in_list = None
        if table:
            out.append(html_table({"table": table, "f": {}}, theme))
            table = []

    for ln in body_md.splitlines():
        s = ln.strip()
        if s.startswith("|"):
            cells = [c.strip() for c in s.strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                table.append(cells)
            continue
        m_ol, m_ul = re.match(r"^\d+[.)]\s+(.*)", s), re.match(r"^[-*]\s+(.*)", s)
        if m_ol or m_ul:
            tag = "ol" if m_ol else "ul"
            if in_list != tag:
                flush()
                out.append("<{}>".format(tag))
                in_list = tag
            out.append("<li>{}</li>".format(inline((m_ol or m_ul).group(1))))
            continue
        flush()
        if not s:
            continue
        h = re.match(r"^(#+)\s+(.*)", s)
        if h:
            lvl = min(len(h.group(1)) + 1, 4)
            out.append("<h{0}>{1}</h{0}>".format(lvl, inline(h.group(2))))
        else:
            out.append("<p>{}</p>".format(inline(s)))
    flush()
    css = (":root{--text:%(text)s;--accent:%(accent)s;--light:%(light)s}" % theme) + """
body{font-family:%s;color:var(--text);background:#fff;margin:0}
main{max-width:760px;margin:40px auto;padding:0 16px;font-size:16px;line-height:1.5}
h2{font-size:20px;border-bottom:3px solid var(--accent);padding-bottom:8px}h3{font-size:16px}
li{margin-bottom:4px}table.data{border-collapse:collapse;width:100%%;font-size:14px}
table.data th,table.data td{border-bottom:1px solid var(--light);padding:4px 8px;text-align:left}
@media print{@page{size:A4;margin:16mm}main{margin:0}}""" % theme["font"]
    return "<!doctype html><html lang=\"en\"><head><meta charset=\"utf-8\"><meta name=\"viewport\" content=\"width=device-width,initial-scale=1\"><title>{}</title><style>{}</style></head><body><main>{}</main></body></html>\n".format(
        esc(meta.get("title", "Memo")), css, "\n".join(out))


# ---------------------------------------------------------------- PPTX

def build_pptx(meta, theme, slides, path):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Emu, Pt

    rgb = lambda c: RGBColor.from_string(c.lstrip("#").upper())
    font = theme["font"].split(",")[0].strip().strip("'\"") or "Arial"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    px = lambda v: Emu(int(v * 12192000 / SW))
    title_only = prs.slide_layouts[5]

    def title_box(sl, x, y, w, h, text, size):
        shp = sl.shapes.title  # real title placeholder, so outline tools and screen readers find it
        shp.left, shp.top, shp.width, shp.height = px(x), px(y), px(w), px(h)
        tf = shp.text_frame
        tf.word_wrap = True
        tf.text = text
        p = tf.paragraphs[0]
        p.alignment = 1  # left
        p.font.size, p.font.bold, p.font.name = Pt(size), True, font
        p.font.color.rgb = rgb(theme["text"])
        tf.vertical_anchor = 1  # top

    def box(sl, x, y, w, h, text, size, bold=False, color=None, items=None, numbered=False):
        tb = sl.shapes.add_textbox(px(x), px(y), px(w), px(h))
        tf = tb.text_frame
        tf.word_wrap = True
        lines = items if items is not None else [text]
        for k, line in enumerate(lines):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            clean = re.sub(r"\*\*(.+?)\*\*", r"\1", line)
            p.text = ("{}. ".format(k + 1) if numbered else ("• " if items is not None else "")) + clean
            p.font.size, p.font.bold, p.font.name = Pt(size), bold, font
            p.font.color.rgb = rgb(color or theme["text"])
            if items is not None:
                p.space_after = Pt(8)
        return tb

    types = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "hbar": XL_CHART_TYPE.BAR_CLUSTERED,
             "line": XL_CHART_TYPE.LINE_MARKERS, "stacked": XL_CHART_TYPE.COLUMN_STACKED,
             "stacked-hbar": XL_CHART_TYPE.BAR_STACKED, "waterfall": XL_CHART_TYPE.COLUMN_STACKED}

    def add_chart(sl, s, x, y, w, h):
        data, ch = chart_data(s), s["f"].get("chart")
        if not data or ch not in types:
            return
        hls = highlights(s)
        cd = CategoryChartData()
        cd.categories = data["cats"]
        if ch == "waterfall":
            steps, _ = waterfall_steps(data)
            cd.add_series("base", [0 if st["total"] else st["lo"] for st in steps])
            cd.add_series(data["series"][0]["name"], [st["hi"] - st["lo"] for st in steps])
        else:
            for ser in data["series"]:
                cd.add_series(ser["name"], [v or 0 for v in ser["vals"]])
        gf = sl.shapes.add_chart(types[ch], px(x), px(y), px(w), px(h), cd)
        chart = gf.chart
        chart.has_title = False
        chart.font.size, chart.font.name = Pt(12), font
        allv = [v for ser in data["series"] for v in ser["vals"] if v is not None]
        numfmt = "#,##0.0" if any(abs(v - round(v)) > 1e-9 for v in allv) else "#,##0"
        chart.font.color.rgb = rgb(theme["text"])
        va = chart.value_axis
        va.has_major_gridlines = False
        if ch != "line":
            va.visible = False
        elif s["f"].get("scale", "zero").lower() != "fit":
            va.minimum_scale = 0
        if ch in ("hbar", "stacked-hbar"):
            try:
                chart.category_axis.reverse_order = True
            except AttributeError:
                pass
        chart.has_legend = len(data["series"]) > 1 and ch != "waterfall"
        if chart.has_legend:
            chart.legend.position, chart.legend.include_in_layout = XL_LEGEND_POSITION.TOP, False
        plot = chart.plots[0]
        if ch != "line":
            plot.gap_width = 60
            if ch in ("stacked", "stacked-hbar", "waterfall"):
                plot.overlap = 100
        cols = series_colors(theme, len(data["series"]), bool(hls))
        if ch == "waterfall":
            base, val = plot.series[0], plot.series[1]
            base.format.fill.background()
            for k, st in enumerate(steps):
                pt = val.points[k]
                pt.format.fill.solid()
                col = theme["accent"] if is_hl(hls, st["cat"], "") else (theme["text"] if st["total"] else theme["muted"])
                pt.format.fill.fore_color.rgb = rgb(col)
                pt.data_label.has_text_frame = True
                pt.data_label.text_frame.text = st["raw"]
                run = pt.data_label.text_frame.paragraphs[0].runs[0]
                run.font.size, run.font.bold = Pt(12), st["total"]
                run.font.color.rgb = rgb(text_on(col))
            return
        plot.has_data_labels = True
        plot.data_labels.font.size = Pt(12)
        plot.data_labels.number_format, plot.data_labels.number_format_is_linked = numfmt, False
        for j, ser in enumerate(plot.series):
            name = data["series"][j]["name"]
            if ch == "line":
                ser.format.line.color.rgb = rgb(theme["accent"] if is_hl(hls, "", name) else cols[j])
                ser.format.line.width = Pt(2.5)
                ser.smooth = False
                continue
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = rgb(cols[j])
            ser.format.line.color.rgb = rgb(theme["background"])
            for k, cat in enumerate(data["cats"]):
                if is_hl(hls, cat, name):
                    pt = ser.points[k]
                    pt.format.fill.solid()
                    pt.format.fill.fore_color.rgb = rgb(theme["accent"])

    for s in slides:
        f, lay = s["f"], s["f"]["layout"]
        sl = prs.slides.add_slide(title_only)
        bg = sl.background.fill
        bg.solid()
        bg.fore_color.rgb = rgb(theme["background"])
        tracker = f.get("tracker") or ("Appendix" if s["id"][0].isalpha() else "")
        if lay == "title":
            box(sl, 96, 180, 1000, 30, meta.get("client", ""), 16, True, theme["accent"])
            title_box(sl, 96, 230, 1000, 200, s["title"], 40)
            if f.get("subtitle"):
                box(sl, 96, 440, 1000, 60, f["subtitle"], 18)
            box(sl, 96, 560, 1000, 30, " | ".join(x for x in (meta.get("audience"), meta.get("date")) if x), 14)
        elif lay == "section":
            box(sl, 96, 240, 1000, 30, tracker, 14, True, theme["accent"])
            title_box(sl, 96, 280, 1000, 200, s["title"], 34)
        else:
            box(sl, 64, 30, 1100, 24, tracker.upper(), 11, True, theme["accent"])
            title_box(sl, 64, 58, 1150, 90, s["title"], 24 if len(s["title"]) < 95 else 21)
            top = 150
            if f.get("header"):
                box(sl, 64, top, 1150, 26, f["header"], 13, color="#5B6475")
                top += 34
            bottom = 640 - (50 if f.get("takeaway") else 0)
            if lay == "summary":
                y = top
                if f.get("lead"):
                    box(sl, 64, y, 1150, 60, f["lead"], 15, color="#3B4458")
                    y += 70
                box(sl, 64, y, 1150, 330, "", 15, items=s["body"], numbered=True)
                if f.get("ask"):
                    box(sl, 64, 570, 1150, 50, f["ask"], 15, True)
            elif lay == "chart":
                w = 720 if s["body"] else 1150
                add_chart(sl, s, 64, top, w, bottom - top)
                if s["body"]:
                    box(sl, 820, top, 400, bottom - top, "", 15, items=s["body"])
                ref = reference(s)
                if ref:
                    box(sl, 64, bottom - 4, 800, 22, "Reference: {} = {}".format(ref[1], fmt(ref[0])), 11, color="#5B6475")
            elif lay == "table" and s["table"]:
                t = s["table"]
                ncol = max(len(r) for r in t)
                w = 720 if s["body"] else 1150
                tbl = sl.shapes.add_table(len(t), ncol, px(64), px(top), px(w), px(min(bottom - top, 34 * len(t)))).table
                for r, row in enumerate(t):
                    for c in range(ncol):
                        cell = tbl.cell(r, c)
                        cell.text = re.sub(r"\*\*(.+?)\*\*", r"\1", row[c] if c < len(row) else "")
                        para = cell.text_frame.paragraphs[0]
                        para.font.size, para.font.name, para.font.bold = Pt(12), font, r == 0
                        para.font.color.rgb = rgb(theme["text"])
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = rgb(theme["light"] if r == 0 else theme["background"])
                if s["body"]:
                    box(sl, 820, top, 400, bottom - top, "", 15, items=s["body"])
            else:
                y = top
                if f.get("lead"):
                    box(sl, 64, y, 1100, 60, f["lead"], 16, color="#3B4458")
                    y += 70
                box(sl, 64, y, 1100, bottom - y, "", 17, items=s["body"])
            if f.get("takeaway"):
                box(sl, 64, 600, 1150, 36, f["takeaway"], 15, True)
        foot = " ".join(x for x in (f.get("footnote", ""), "Source: " + f["source"] if f.get("source") else "") if x)
        if foot:
            box(sl, 64, 660, 1050, 40, foot, 10, color="#5B6475")
        box(sl, 1160, 668, 60, 24, s["id"], 11, True)
        if f.get("note"):
            sl.notes_slide.notes_text_frame.text = f["note"]
    prs.save(str(path))


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description="Render a Markdown deck spec to HTML (always) and PPTX (if python-pptx is installed).")
    ap.add_argument("spec", help="path to the deck spec, e.g. engagement/<client>/06-deck.md")
    ap.add_argument("--out-dir", help="output folder (default: next to the spec)")
    ap.add_argument("--name", help="output file stem (default: spec file stem)")
    ap.add_argument("--no-pptx", action="store_true", help="skip the .pptx even if python-pptx is installed")
    ap.add_argument("--titles", action="store_true", help="print the titles-only read-through and exit")
    ap.add_argument("--check", action="store_true", help="print spec warnings and exit without writing files")
    args = ap.parse_args()

    spec = Path(args.spec)
    if not spec.is_file():
        sys.exit("spec not found: {}".format(spec))
    text = spec.read_text(encoding="utf-8")
    meta, theme, slides = parse_spec(text)
    memo = meta.get("format", "deck").lower() == "memo"

    if args.titles:
        for s in slides:
            print("{:>4}  {}".format(s["id"], s["title"]))
        return
    if not memo and not slides:
        sys.exit("no slides found: each slide starts with a heading like '## 3. <action title>'")

    warns = [] if memo else lint_slides(meta, slides)
    for w in warns:
        print("WARNING " + w)
    if args.check:
        print("{} slides, {} warnings".format(len(slides), len(warns)))
        return

    out_dir = Path(args.out_dir) if args.out_dir else spec.parent
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = args.name or spec.stem
    html_path = out_dir / (stem + ".html")
    html_path.write_text(build_memo_html(meta, theme, text) if memo else build_html(meta, theme, slides), encoding="utf-8")
    print("wrote {} ({} slides)".format(html_path, len(slides)) if not memo else "wrote {} (memo)".format(html_path))

    if memo or args.no_pptx:
        return
    try:
        import pptx  # noqa: F401
    except ImportError:
        print("python-pptx not installed, skipped .pptx. Install with: python3 -m pip install python-pptx")
        return
    pptx_path = out_dir / (stem + ".pptx")
    build_pptx(meta, theme, slides, pptx_path)
    print("wrote {}".format(pptx_path))


if __name__ == "__main__":
    main()
